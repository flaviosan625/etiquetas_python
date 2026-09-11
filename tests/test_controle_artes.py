import pathlib
import sys

sys.path.insert(0, str(pathlib.Path(__file__).resolve().parent.parent))

import openpyxl

import controle_artes as ct

CADERNO = "CAD ARTES EXECUTIVO"


def _ficha(nome, medidas, situacao="APROVADO", slide=7, material="LONA IMPRESSA",
           sangria="", qtd="01", links=None, nome_arquivo=""):
    return {"slide": slide, "nome": nome, "material": material, "medidas": medidas,
            "medidas_sangria": sangria, "quantidade": qtd, "situacao": situacao,
            "links": links or [], "nome_arquivo": nome_arquivo, "motivo_nome": None}


def _item(nome, largura, altura, status="APROVADO", area="PALCO ARENA",
          material="LONA", quantidade=1):
    return {"area": area, "nome": nome, "quantidade": quantidade,
            "largura": largura, "altura": altura, "material": material,
            "status": status}


def _planilha_do_cliente(caminho, linhas):
    """
    Reproduz o formato do cliente: primeira coluna vazia, cabecalho na
    linha 2, e a AREA em celula mesclada (so a primeira linha do bloco
    traz o nome).
    """
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = ct.ABA_CLIENTE
    ws.cell(row=2, column=2, value="ÁREA")
    ws.cell(row=2, column=3, value="NOME")
    ws.cell(row=2, column=9, value="STATUS")
    for i, (area, nome, largura, altura, material, status) in enumerate(linhas, start=3):
        if area:
            ws.cell(row=i, column=2, value=area)
        ws.cell(row=i, column=3, value=nome)
        ws.cell(row=i, column=4, value=1)
        ws.cell(row=i, column=5, value=largura)
        ws.cell(row=i, column=6, value=altura)
        ws.cell(row=i, column=7, value=material)
        ws.cell(row=i, column=9, value=status)
    wb.save(str(caminho))
    wb.close()
    return caminho


# ------------------------------------------------------- nada para tras

def test_peca_que_so_existe_no_caderno_fica_marcada():
    linhas = ct.montar_linhas([_ficha("LONA FRONTAL PORTICO", "12,80 X 4,50")], [], CADERNO)

    assert len(linhas) == 1
    assert linhas[0]["confere"] == ct.MARCA_SO_CADERNO


def test_peca_que_so_existe_na_planilha_do_cliente_nao_se_perde():
    """
    O motivo de esta planilha existir. No primeiro cruzamento real foram
    63 pecas cobradas pelo cliente sem slide nenhum do nosso lado — e
    isso era invisivel olhando os dois documentos separados.
    """
    linhas = ct.montar_linhas([], [_item("PAREDE CURVA", 21.3, 6.0)], CADERNO)

    assert len(linhas) == 1
    assert linhas[0]["confere"] == ct.MARCA_SO_PLANILHA
    assert linhas[0]["peca"] == "PAREDE CURVA"


def test_par_exige_nome_parecido_alem_da_medida_igual():
    """
    Medida sozinha colide: 5,00 x 3,00 casou 'PAINEL DE FUNDO' com
    'VISTA PAREDE DO DEPOSITO', de outra area (2026-09-11). Melhor dizer
    'sem par' do que inventar um.
    """
    linhas = ct.montar_linhas(
        [_ficha("PAINEL DE FUNDO", "5,00 X 3,00")],
        [_item("VISTA PAREDE DO DEPOSITO", 5.0, 3.0, area="SALA CONEXOES")],
        CADERNO)

    so_caderno = [l for l in linhas if l["origem"] == "caderno"]
    assert so_caderno[0]["confere"] == ct.MARCA_SO_CADERNO
    assert any(l["confere"] == ct.MARCA_SO_PLANILHA for l in linhas), \
        "o item da planilha continua na lista, nao some junto"


def test_nomes_diferentes_com_parentesco_e_medida_igual_casam():
    """A mesma peca e 'TESTEIRA ARENA' no caderno e 'TESTEIRA - LONA' na
    planilha — caso real conferido em 2026-09-11."""
    linhas = ct.montar_linhas(
        [_ficha("TESTEIRA ARENA", "22,17 X 2,00")],
        [_item("TESTEIRA - LONA", 22.17, 2.0)],
        CADERNO)

    assert len(linhas) == 1, "casou: nao duplica a peca"
    assert linhas[0]["confere"] == ct.MARCA_CONFERE
    assert linhas[0]["status_planilha"] == "APROVADO"


def test_status_discordando_entre_os_dois_documentos_vira_divergente():
    linhas = ct.montar_linhas(
        [_ficha("TESTEIRA ARENA", "22,17 X 2,00", situacao="EM CRIACAO")],
        [_item("TESTEIRA - LONA", 22.17, 2.0, status="APROVADO")],
        CADERNO)

    assert linhas[0]["confere"] == ct.MARCA_DIVERGENTE


# ---------------------------------------------------- observacao do uso

def test_observacao_sobrevive_a_peca_renomeada_pelo_cliente():
    """Renomeou a peca, mas o slide e o mesmo: a ancora do slide salva."""
    antes = ct.montar_linhas([_ficha("LONA FRONTAL PORTICO", "12,80 X 4,50", slide=7)], [], CADERNO)
    guardadas = [{"ancoras": antes[0]["_ancoras"], "texto": "FALAR COM A JULIANA"}]

    depois = ct.montar_linhas(
        [_ficha("LONA FRONTAL DO PORTICO EXTERNO", "12,80 X 4,50", slide=7)],
        [], CADERNO, guardadas)

    assert depois[0]["observacao"] == "FALAR COM A JULIANA"
    assert not any(l["confere"] == ct.MARCA_ORFA for l in depois)


def test_observacao_sobrevive_a_slide_movido_de_lugar():
    """Inseriram slide antes: o numero mudou, mas o nome da peca salva."""
    antes = ct.montar_linhas([_ficha("LONA FRONTAL PORTICO", "12,80 X 4,50", slide=7)], [], CADERNO)
    guardadas = [{"ancoras": antes[0]["_ancoras"], "texto": "CONFERIR SANGRIA"}]

    depois = ct.montar_linhas(
        [_ficha("LONA FRONTAL PORTICO", "12,80 X 4,50", slide=9)],
        [], CADERNO, guardadas)

    assert depois[0]["observacao"] == "CONFERIR SANGRIA"


def test_observacao_que_perdeu_as_duas_ancoras_nao_e_apagada():
    """Apagar em silencio seria perder trabalho de gente."""
    depois = ct.montar_linhas(
        [_ficha("OUTRA PECA QUALQUER", "1,00 X 1,00", slide=3)],
        [], CADERNO,
        [{"ancoras": ["CAD ARTES EXECUTIVO|slide|7"],
          "texto": "ESSA PECA SUMIU DO CADERNO"}])

    orfas = [l for l in depois if l["confere"] == ct.MARCA_ORFA]
    assert len(orfas) == 1
    assert orfas[0]["observacao"] == "ESSA PECA SUMIU DO CADERNO"


# ------------------------------------------------ planilha real do fluxo

def test_area_em_celula_mesclada_e_arrastada_para_as_linhas_de_baixo(tmp_path):
    """Sem isso, 58 das 63 pecas ficariam sem area."""
    caminho = _planilha_do_cliente(tmp_path / "status.xlsx", [
        ("PALCO ARENA", "P1 - ESQUERDA FRENTE", 2.81, 5.5, "LONA", "EM APROVAÇÃO"),
        (None, "P1 - ESQUERDA VERSO", 13.18, 5.5, "LONA", "EM APROVAÇÃO"),
        ("SALA CONEXOES", "VISTA 01 - EXTERNA", 3.0, 2.0, "LONA", "APROVADO"),
    ])

    itens = ct.ler_planilha_do_cliente(caminho)

    assert [i["area"] for i in itens] == ["PALCO ARENA", "PALCO ARENA", "SALA CONEXOES"]


def test_planilha_de_controle_vai_e_volta_com_a_observacao(tmp_path):
    destino = tmp_path / "CONTROLE.xlsx"
    fichas = [_ficha("LONA FRONTAL PORTICO", "12,80 X 4,50",
                     nome_arquivo="1UN LONA 12,80X4,50M_FRONTAL PORTICO")]

    linhas = ct.montar_linhas(fichas, [], CADERNO)
    ct.escrever(destino, linhas, "CONTROLE DE TESTE")

    wb = openpyxl.load_workbook(str(destino))
    ws = wb[ct.ABA]
    coluna = ct._CAMPOS.index("observacao") + 1
    ws.cell(row=ct.LINHA_CABECALHO + 1, column=coluna, value="COMPRAR LONA FOSCA")
    wb.save(str(destino))
    wb.close()

    guardadas = ct.ler_observacoes(destino)
    refeitas = ct.montar_linhas(fichas, [], CADERNO, guardadas)

    assert refeitas[0]["observacao"] == "COMPRAR LONA FOSCA"


def test_ler_observacoes_de_arquivo_que_nao_existe_ou_esta_quebrado(tmp_path):
    """Perder observacao e ruim; derrubar a geracao por causa dela e pior."""
    assert ct.ler_observacoes(tmp_path / "nem_existe.xlsx") == []

    quebrado = tmp_path / "quebrado.xlsx"
    quebrado.write_bytes(b"isto nao e uma planilha")
    assert ct.ler_observacoes(quebrado) == []
